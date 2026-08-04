# Copyright 2025 the LlamaFactory team.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""使用中医疾病、中医证候和 ICD-10 词表归一化 AI 诊断标签."""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import sys
import tempfile
from collections import Counter, defaultdict
from collections.abc import Callable, Iterator
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]
if __package__ in (None, ""):
    sys.path.insert(0, str(PROJECT_ROOT))

from medical.analysis.engine import (  # noqa: E402
    DIAGNOSIS_SEPARATOR_PATTERN,
    normalize_diagnosis_value,
    normalize_sickness_value,
    normalize_syndrome_value,
)


DEFAULT_INPUT = PROJECT_ROOT / "medical/processed_data/doctor_43_朱子奇/medical_records_ai.jsonl"
DEFAULT_TCM_DISEASE_MAPPING = PROJECT_ROOT / "medical/data/辨病.xlsx"
DEFAULT_TCM_SYNDROME_MAPPING = PROJECT_ROOT / "medical/data/辩证.xlsx"
DEFAULT_WESTERN_DIAGNOSIS_MAPPING = PROJECT_ROOT / "medical/data/ICD10.xlsx"


@dataclass(frozen=True)
class MappingSpec:
    field_name: str
    display_name: str
    workbook_path: Path
    label_header: str
    normalizer: Callable[[Any], str]
    aliases_share_code: bool


@dataclass
class LabelMapping:
    spec: MappingSpec
    alias_to_standard: dict[str, str]
    normalized_to_standard: dict[str, str]
    standard_names: tuple[str, ...]
    ambiguous_aliases: dict[str, tuple[str, ...]]
    ambiguous_normalized_names: dict[str, tuple[str, ...]]


@dataclass
class FieldStats:
    nonempty_records: int = 0
    changed_records: int = 0
    fully_mapped_records: int = 0
    partially_mapped_records: int = 0
    unmapped_records: int = 0
    mapped_label_occurrences: int = 0
    unmatched: Counter[str] = field(default_factory=Counter)


def _clean_cell(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _unique_mappings(candidates: dict[str, set[str]]) -> dict[str, str]:
    return {name: next(iter(standards)) for name, standards in candidates.items() if len(standards) == 1}


def load_label_mapping(spec: MappingSpec) -> LabelMapping:
    """读取标准名和别名，建立名称→标准名映射，不将编码写入结果."""
    if not spec.workbook_path.is_file():
        raise FileNotFoundError(f"映射文件不存在: {spec.workbook_path}")
    workbook = load_workbook(spec.workbook_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = tuple(_clean_cell(value) for value in next(rows, ()))
        if "编码" not in headers or spec.label_header not in headers:
            raise ValueError(f"{spec.workbook_path} 表头必须包含 编码 和 {spec.label_header}，实际为: {headers}")
        code_index = headers.index("编码")
        label_index = headers.index(spec.label_header)
        names_by_code: dict[str, list[str]] = defaultdict(list)
        for row_number, row in enumerate(rows, start=2):
            code = _clean_cell(row[code_index] if code_index < len(row) else None)
            label = _clean_cell(row[label_index] if label_index < len(row) else None)
            if not code and not label:
                continue
            if not code or not label:
                raise ValueError(f"{spec.workbook_path} 第 {row_number} 行编码或标签为空")
            if label not in names_by_code[code]:
                names_by_code[code].append(label)
        if not names_by_code:
            raise ValueError(f"映射文件没有有效标签: {spec.workbook_path}")

        alias_candidates: dict[str, set[str]] = defaultdict(set)
        normalized_candidates: dict[str, set[str]] = defaultdict(set)
        standard_names = []
        for names in names_by_code.values():
            groups = [(names[0], names)] if spec.aliases_share_code else [(name, [name]) for name in names]
            for standard_name, aliases in groups:
                standard_names.append(standard_name)
                for alias in aliases:
                    alias_candidates[alias].add(standard_name)
                    normalized_alias = spec.normalizer(alias)
                    if normalized_alias:
                        normalized_candidates[normalized_alias].add(standard_name)

        return LabelMapping(
            spec=spec,
            alias_to_standard=_unique_mappings(alias_candidates),
            normalized_to_standard=_unique_mappings(normalized_candidates),
            standard_names=tuple(dict.fromkeys(standard_names)),
            ambiguous_aliases={
                alias: tuple(sorted(standards)) for alias, standards in alias_candidates.items() if len(standards) > 1
            },
            ambiguous_normalized_names={
                name: tuple(sorted(standards))
                for name, standards in normalized_candidates.items()
                if len(standards) > 1
            },
        )
    finally:
        workbook.close()


def normalize_label_value(value: Any, mapping: LabelMapping) -> tuple[Any, int, list[str]]:
    """按顺序拆分多诊断，将标准名或别名统一替换为标准名."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return value, 0, []
    if not isinstance(value, str):
        return value, 0, [_clean_cell(value)]
    if not mapping.spec.normalizer(value):
        return value, 0, [_clean_cell(value)]

    raw_parts = [part.strip() for part in re.split(DIAGNOSIS_SEPARATOR_PATTERN, value) if part.strip()]
    if not raw_parts:
        return value, 0, [_clean_cell(value)]
    mapped_parts = []
    matched_count = 0
    unmatched = []
    for raw_part in raw_parts:
        standard_name = mapping.alias_to_standard.get(raw_part)
        if standard_name is not None:
            mapped_parts.append(standard_name)
            matched_count += 1
            continue

        normalized = mapping.spec.normalizer(raw_part)
        normalized_parts = [part for part in normalized.split("、") if part]
        standards = [mapping.normalized_to_standard.get(part) for part in normalized_parts]
        if standards and all(standards):
            mapped_parts.extend(standards)
            matched_count += len(standards)
        else:
            mapped_parts.append(raw_part)
            unmatched.append(raw_part)
    return "、".join(dict.fromkeys(mapped_parts)), matched_count, unmatched


def iter_jsonl(path: Path) -> Iterator[tuple[int, dict[str, Any]]]:
    with path.open(encoding="utf-8") as source:
        for line_number, line in enumerate(source, start=1):
            if not line.strip():
                continue
            try:
                row = json.loads(line)
            except json.JSONDecodeError as error:
                raise ValueError(f"{path} 第 {line_number} 行不是合法 JSON: {error}") from error
            if not isinstance(row, dict):
                raise ValueError(f"{path} 第 {line_number} 行必须是 JSON 对象")
            yield line_number, row


def _mapping_report(mapping: LabelMapping) -> dict[str, Any]:
    return {
        "workbook": str(mapping.spec.workbook_path),
        "standard_name_count": len(mapping.standard_names),
        "alias_name_count": len(mapping.alias_to_standard),
        "normalized_name_count": len(mapping.normalized_to_standard),
        "ambiguous_aliases": {name: list(values) for name, values in mapping.ambiguous_aliases.items()},
        "ambiguous_normalized_names": {
            name: list(values) for name, values in mapping.ambiguous_normalized_names.items()
        },
    }


def _stats_report(stats: FieldStats) -> dict[str, Any]:
    return {
        "nonempty_records": stats.nonempty_records,
        "changed_records": stats.changed_records,
        "fully_mapped_records": stats.fully_mapped_records,
        "partially_mapped_records": stats.partially_mapped_records,
        "unmapped_records": stats.unmapped_records,
        "mapped_label_occurrences": stats.mapped_label_occurrences,
        "unmatched_label_occurrences": sum(stats.unmatched.values()),
        "top_unmatched_labels": [
            {"label": label, "count": count} for label, count in stats.unmatched.most_common(100)
        ],
    }


def normalize_jsonl(
    input_path: Path,
    mappings: list[LabelMapping],
    *,
    output_path: Path | None,
) -> dict[str, Any]:
    """流式归一化 JSONL；output_path 为空时只统计，不写文件."""
    if not input_path.is_file():
        raise FileNotFoundError(f"输入文件不存在: {input_path}")
    stats = {mapping.spec.field_name: FieldStats() for mapping in mappings}
    record_count = 0
    temporary_path: Path | None = None
    destination = None
    if output_path is not None:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        temporary = tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            dir=output_path.parent,
            prefix=f".{output_path.name}.",
            suffix=".tmp",
            delete=False,
        )
        temporary_path = Path(temporary.name)
        destination = temporary
    try:
        for _, row in iter_jsonl(input_path):
            record_count += 1
            for mapping in mappings:
                field_name = mapping.spec.field_name
                old_value = row.get(field_name)
                order_sn = row.get("order_sn")
                if old_value is None or (isinstance(old_value, str) and not old_value.strip()):
                    continue
                field_stats = stats[field_name]
                field_stats.nonempty_records += 1
                # 归一化后的诊断名称。成功匹配并归一化的诊断项数量。未能匹配标准名称的诊断项列表, old_value 为归一化前
                new_value, matched_count, unmatched = normalize_label_value(old_value, mapping)
                print(order_sn, field_name, old_value, new_value, matched_count, unmatched)
                # import pdb
                # pdb.set_trace()
                field_stats.mapped_label_occurrences += matched_count
                field_stats.unmatched.update(unmatched)
                if matched_count and not unmatched:
                    field_stats.fully_mapped_records += 1
                elif matched_count:
                    field_stats.partially_mapped_records += 1
                else:
                    field_stats.unmapped_records += 1
                if new_value != old_value:
                    row[field_name] = new_value
                    field_stats.changed_records += 1
            if destination is not None:
                destination.write(json.dumps(row, ensure_ascii=False, separators=(",", ":")) + "\n")
        if destination is not None:
            destination.flush()
            os.fsync(destination.fileno())
            destination.close()
            os.replace(temporary_path, output_path)
            temporary_path = None
    finally:
        if destination is not None and not destination.closed:
            destination.close()
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
    return {
        "input": str(input_path),
        "output": str(output_path) if output_path else None,
        "record_count": record_count,
        "mappings": {mapping.spec.field_name: _mapping_report(mapping) for mapping in mappings},
        "fields": {field_name: _stats_report(field_stats) for field_name, field_stats in stats.items()},
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="使用三份标准词表归一化 JSONL 中的 AI 诊断标签")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--tcm-disease-mapping", type=Path, default=DEFAULT_TCM_DISEASE_MAPPING)
    parser.add_argument("--tcm-syndrome-mapping", type=Path, default=DEFAULT_TCM_SYNDROME_MAPPING)
    parser.add_argument("--western-diagnosis-mapping", type=Path, default=DEFAULT_WESTERN_DIAGNOSIS_MAPPING)
    parser.add_argument("--output", type=Path, help="输出 JSONL；默认在输入文件旁生成 *_normalized.jsonl")
    parser.add_argument("--audit-output", type=Path, help="审计 JSON 输出路径；默认与输出 JSONL 同名并加 .audit.json")
    parser.add_argument("--check-only", action="store_true", help="只统计映射结果，不写 JSONL")
    parser.add_argument("--in-place", action="store_true", help="原地替换输入文件，并先创建不可覆盖的备份")
    parser.add_argument("--overwrite", action="store_true", help="允许覆盖已存在的非输入输出文件")
    return parser


def _mapping_specs(args: argparse.Namespace) -> list[MappingSpec]:
    return [
        MappingSpec(
            "ai_diagnosis_sickness",
            "中医疾病",
            args.tcm_disease_mapping,
            "辨病",
            normalize_sickness_value,
            True,
        ),
        MappingSpec(
            "ai_diagnosis_disease",
            "中医证候",
            args.tcm_syndrome_mapping,
            "辩证",
            normalize_syndrome_value,
            True,
        ),
        MappingSpec(
            "ai_diagnosis_illness",
            "西医诊断",
            args.western_diagnosis_mapping,
            "ICD10",
            normalize_diagnosis_value,
            False,
        ),
    ]


@lru_cache(maxsize=1)
def load_default_label_mappings() -> dict[str, LabelMapping]:
    """按需加载三份默认词表，同一进程内只读取一次 Excel."""
    args = argparse.Namespace(
        tcm_disease_mapping=DEFAULT_TCM_DISEASE_MAPPING,
        tcm_syndrome_mapping=DEFAULT_TCM_SYNDROME_MAPPING,
        western_diagnosis_mapping=DEFAULT_WESTERN_DIAGNOSIS_MAPPING,
    )
    mappings = [load_label_mapping(spec) for spec in _mapping_specs(args)]
    return {mapping.spec.field_name: mapping for mapping in mappings}


def normalize_ai_diagnosis_to_standard(field_name: str, value: Any) -> Any:
    """先执行字段现有别名规则，再通过 Excel mapping 替换为标准名."""
    mappings = load_default_label_mappings()
    if field_name not in mappings:
        raise ValueError(f"不支持的 AI 诊断字段: {field_name}")
    mapping = mappings[field_name]
    alias_normalized = mapping.spec.normalizer(value)
    if not alias_normalized:
        return alias_normalized
    standard_value, _, _ = normalize_label_value(alias_normalized, mapping)
    return standard_value


def _default_output(input_path: Path) -> Path:
    return input_path.with_name(f"{input_path.stem}_normalized{input_path.suffix}")


def main() -> None:
    args = build_parser().parse_args()
    if args.check_only and args.in_place:
        raise SystemExit("--check-only 与 --in-place 不能同时使用")
    if args.in_place and args.output:
        raise SystemExit("--in-place 与 --output 不能同时使用")

    input_path = args.input.expanduser().resolve()
    mappings = [load_label_mapping(spec) for spec in _mapping_specs(args)]
    output_path = (
        None if args.check_only else (input_path if args.in_place else (args.output or _default_output(input_path)))
    )
    if output_path is not None:
        output_path = output_path.expanduser().resolve()
        if output_path != input_path and output_path.exists() and not args.overwrite:
            raise SystemExit(f"输出文件已存在，使用 --overwrite 允许覆盖: {output_path}")

    audit_output = args.audit_output
    if audit_output is None and output_path is not None:
        audit_output = output_path.with_suffix(output_path.suffix + ".audit.json")
    if audit_output is not None:
        audit_output = audit_output.expanduser().resolve()
        protected_paths = {input_path}
        if output_path is not None:
            protected_paths.add(output_path)
        if audit_output in protected_paths:
            raise SystemExit(f"审计文件不能与输入或输出 JSONL 同路径: {audit_output}")
        if audit_output.exists() and not args.overwrite:
            raise SystemExit(f"审计文件已存在，使用 --overwrite 允许覆盖: {audit_output}")

    backup_path = None
    if args.in_place:
        backup_path = input_path.with_name(f"{input_path.stem}.before_diagnosis_normalization{input_path.suffix}")
        if backup_path.exists():
            raise SystemExit(f"备份文件已存在，拒绝覆盖: {backup_path}")
        shutil.copy2(input_path, backup_path)

    report = normalize_jsonl(input_path, mappings, output_path=output_path)
    if backup_path is not None:
        report["backup"] = str(backup_path)

    if audit_output is not None:
        audit_output.parent.mkdir(parents=True, exist_ok=True)
        audit_output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
