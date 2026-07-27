# Copyright 2026 the LlamaFactory team.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Export normalized diagnosis fields to one Excel sheet per online doctor."""

# ruff: noqa: E402

import argparse
import json
import math
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


PROJECT_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT_ROOT))

from medical.analysis.engine import normalize_diagnosis_value, normalize_syndrome_value


DEFAULT_DATA_DIR = PROJECT_ROOT / "medical" / "data"
DEFAULT_OUTPUT_FILE = Path(__file__).resolve().parent / "output" / "normalized_diagnoses_by_doctor.xlsx"
ONLINE_FILE_MARKER = "_AI医生分身混合问诊数据_"
HEADERS = (
    "序号",
    "order_sn",
    "diagnosis_illness（归一化）",
    "diagnosis_sickness（归一化）",
    "diagnosis_disease（归一化）",
)


def discover_online_files(data_dir: Path) -> list[Path]:
    """Find JSON exports in online_* directories and online export files at the data root."""
    files = []
    for path in data_dir.rglob("*.json"):
        relative_parts = path.relative_to(data_dir).parts
        if any(part.startswith(".") or part == "__MACOSX" for part in relative_parts):
            continue
        in_online_directory = any(part.startswith("online_") for part in relative_parts[:-1])
        if in_online_directory or ONLINE_FILE_MARKER in path.name:
            files.append(path)
    return sorted(set(files))


def doctor_name_from_file(path: Path) -> str:
    """Infer the doctor name when a source record does not contain doctor_name."""
    if ONLINE_FILE_MARKER in path.name:
        return path.name.split(ONLINE_FILE_MARKER, maxsplit=1)[0].strip()
    return path.parent.name.removeprefix("online_").strip()


def load_records(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8") as file:
        records = json.load(file)
    if not isinstance(records, list):
        raise ValueError(f"Online doctor data must be a JSON list: {path}")
    return [record for record in records if isinstance(record, dict)]


def normalize_record(record: dict[str, Any]) -> dict[str, str]:
    """Return the three diagnosis fields after applying the analysis normalization rules."""
    return {
        "order_sn": str(record.get("order_sn") or "").strip(),
        "diagnosis_illness": normalize_diagnosis_value(record.get("diagnosis_illness")),
        "diagnosis_sickness": normalize_diagnosis_value(record.get("diagnosis_sickness")),
        "diagnosis_disease": normalize_syndrome_value(record.get("diagnosis_disease")),
    }


def collect_doctor_records(paths: list[Path]) -> dict[str, list[dict[str, str]]]:
    """Group normalized encounter records by doctor."""
    grouped: dict[str, list[dict[str, str]]] = defaultdict(list)
    for path in paths:
        fallback_doctor = doctor_name_from_file(path)
        for record in load_records(path):
            doctor = str(record.get("doctor_name") or fallback_doctor).strip() or "未知医生"
            grouped[doctor].append(normalize_record(record))
    return dict(sorted(grouped.items()))


def safe_excel_text(value: Any) -> str:
    """Keep identifiers and clinical text from being interpreted as Excel formulas."""
    text = str(value or "")
    return f"'{text}" if text.startswith("=") else text


def safe_sheet_title(value: str, used_titles: set[str]) -> str:
    """Create a unique Excel-compatible sheet name."""
    base = re.sub(r"[\[\]:*?/\\]", "_", value).strip()[:31] or "未知医生"
    title = base
    index = 2
    while title in used_titles:
        suffix = f"_{index}"
        title = f"{base[: 31 - len(suffix)]}{suffix}"
        index += 1
    used_titles.add(title)
    return title


def _column_width(records: list[dict[str, str]], field: str, minimum: int, maximum: int) -> int:
    longest = max((len(record[field]) for record in records), default=minimum)
    return min(maximum, max(minimum, longest + 2))


def _write_doctor_sheet(
    workbook: Workbook,
    doctor: str,
    records: list[dict[str, str]],
    sheet_index: int,
    used_titles: set[str],
) -> None:
    sheet = workbook.create_sheet(safe_sheet_title(doctor, used_titles))
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append(HEADERS)

    for index, record in enumerate(records, start=1):
        sheet.append(
            (
                index,
                safe_excel_text(record["order_sn"]),
                safe_excel_text(record["diagnosis_illness"]),
                safe_excel_text(record["diagnosis_sickness"]),
                safe_excel_text(record["diagnosis_disease"]),
            )
        )

    header_fill = PatternFill("solid", fgColor="1D6B50")
    header_font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Microsoft YaHei", size=10, color="16231D")
    row_border = Border(bottom=Side(style="thin", color="E5E0D5"))
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 28

    diagnosis_widths = (
        _column_width(records, "diagnosis_illness", 28, 58),
        _column_width(records, "diagnosis_sickness", 28, 48),
        _column_width(records, "diagnosis_disease", 28, 58),
    )
    for row in sheet.iter_rows(min_row=2):
        row[0].alignment = Alignment(horizontal="right", vertical="center")
        row[1].number_format = "@"
        for cell in row:
            cell.font = body_font
            cell.border = row_border
            cell.alignment = Alignment(
                horizontal=cell.alignment.horizontal or "left",
                vertical="center",
                wrap_text=cell.column >= 3,
            )
        required_lines = max(
            math.ceil(len(str(row[column].value or "")) / diagnosis_widths[column - 2]) for column in range(2, 5)
        )
        sheet.row_dimensions[row[0].row].height = min(60, max(20, required_lines * 16))

    sheet.column_dimensions["A"].width = 9
    sheet.column_dimensions["B"].width = _column_width(records, "order_sn", 20, 28)
    sheet.column_dimensions["C"].width = diagnosis_widths[0]
    sheet.column_dimensions["D"].width = diagnosis_widths[1]
    sheet.column_dimensions["E"].width = diagnosis_widths[2]

    last_row = max(sheet.max_row, 1)
    table = Table(displayName=f"NormalizedDiagnoses{sheet_index}", ref=f"A1:E{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium4",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def export_normalized_diagnoses(data_dir: Path, output_file: Path) -> dict[str, int]:
    """Generate the normalized diagnosis workbook and return record counts by doctor."""
    source_files = discover_online_files(data_dir)
    if not source_files:
        raise FileNotFoundError(f"No online doctor JSON files found under: {data_dir}")

    grouped = collect_doctor_records(source_files)
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_titles: set[str] = set()
    for sheet_index, (doctor, records) in enumerate(grouped.items(), start=1):
        _write_doctor_sheet(workbook, doctor, records, sheet_index, used_titles)

    output_file.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_file)
    return {doctor: len(records) for doctor, records in grouped.items()}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="按医生导出三个归一化诊断字段，每位医生一个 Excel Sheet。")
    parser.add_argument("--data-dir", type=Path, default=DEFAULT_DATA_DIR, help=f"数据目录，默认：{DEFAULT_DATA_DIR}")
    parser.add_argument(
        "--output", type=Path, default=DEFAULT_OUTPUT_FILE, help=f"输出文件，默认：{DEFAULT_OUTPUT_FILE}"
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    counts = export_normalized_diagnoses(args.data_dir, args.output)
    print(f"Exported {sum(counts.values())} records for {len(counts)} doctors to {args.output}")
    print(json.dumps(counts, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
