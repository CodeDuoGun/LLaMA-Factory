# Copyright 2026 the LlamaFactory team.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


sys.path.insert(0, str(Path(__file__).resolve().parents[3]))

from medical.analysis.export_normalized_diagnoses import export_normalized_diagnoses  # noqa: E402


def _write_records(path: Path, doctor: str, records: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    for record in records:
        record["doctor_name"] = doctor
    path.write_text(json.dumps(records, ensure_ascii=False), encoding="utf-8")


def test_export_normalizes_three_fields_and_creates_one_sheet_per_doctor(tmp_path: Path) -> None:
    data_dir = tmp_path / "data"
    _write_records(
        data_dir / "online_a" / "甲医生_AI医生分身混合问诊数据_2_test.json",
        "甲医生",
        [
            {
                "order_sn": "ORDER-1",
                "diagnosis_illness": "肺结节病",
                "diagnosis_sickness": "咳嗽病",
                "diagnosis_disease": "肝气不舒",
            },
            {
                "order_sn": "ORDER-2",
                "diagnosis_illness": "十二指肠球部溃疡,慢性菱缩性胃炎",
                "diagnosis_sickness": "",
                "diagnosis_disease": "脾肺气虚，痰瘀互结证",
            },
        ],
    )
    _write_records(
        data_dir / "online_b" / "乙医生_AI医生分身混合问诊数据_1_test.json",
        "乙医生",
        [
            {
                "order_sn": "ORDER-3",
                "diagnosis_illness": "双肺间质性改变伴纤维",
                "diagnosis_sickness": "鼻鼽病",
                "diagnosis_disease": "痰瘀互结证/脾肺气虚",
            }
        ],
    )

    output = tmp_path / "normalized.xlsx"
    counts = export_normalized_diagnoses(data_dir, output)
    workbook = load_workbook(output)

    assert counts == {"乙医生": 1, "甲医生": 2}
    assert workbook.sheetnames == ["乙医生", "甲医生"]
    assert [cell.value for cell in workbook["甲医生"][1]] == [
        "序号",
        "order_sn",
        "diagnosis_illness（归一化）",
        "diagnosis_sickness（归一化）",
        "diagnosis_disease（归一化）",
    ]
    assert [cell.value for cell in workbook["甲医生"][2]] == [
        1,
        "ORDER-1",
        "肺结节",
        "咳嗽",
        "肝气不舒证",
    ]
    assert [cell.value for cell in workbook["甲医生"][3]] == [
        2,
        "ORDER-2",
        "十二指肠球部溃疡、慢性萎缩性胃炎",
        None,
        "痰瘀互结证、脾肺气虚证",
    ]
    assert [cell.value for cell in workbook["乙医生"][2]] == [
        1,
        "ORDER-3",
        "双肺间质性改变伴纤维化",
        "鼻鼽病",
        "痰瘀互结证、脾肺气虚证",
    ]
    assert workbook["甲医生"].freeze_panes == "A2"
    assert len(workbook["甲医生"].tables) == 1
